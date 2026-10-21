"""Editable Word version of a meeting: summary and action items first, transcript after, speakers in bold."""
import io
import os
import re

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import tfsfmt

LOGO = "/opt/tfs-review/static/tfs-logo.png"
GREEN, TERRA, INK, MUT = RGBColor(0x4A, 0x5E, 0x3D), RGBColor(0xC0, 0x56, 0x2B), RGBColor(0x2E, 0x21, 0x18), RGBColor(0x6B, 0x5A, 0x4B)


def _rgb(hexs: str) -> RGBColor:
    h = hexs.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _runs(par, text: str, size=None, color=None):
    """Add text to a paragraph, turning **bold** into bold runs."""
    for i, piece in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not piece:
            continue
        r = par.add_run(piece)
        r.bold = (i % 2 == 1)
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color


def _style(doc, name, base="Normal", font="Calibri", size=11, bold=False, color=INK, before=0, after=6):
    from docx.enum.style import WD_STYLE_TYPE
    try:
        st = doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles[base]
    st.font.name = font
    st.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    return st


def _rule(par, hexcolor="C0562B"):
    p = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "10"), ("w:space", "4"), ("w:color", hexcolor)):
        bt.set(qn(k), v)
    b.append(bt)
    p.append(b)


def make_docx(title: str, when: str, summary: str, transcript: str, note: str = "", info: str = "", people=None) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin, sec.bottom_margin = Inches(0.8), Inches(0.8)
    _style(doc, "Normal", size=11, after=6)
    _style(doc, "Title", font="Georgia", size=26, bold=True, color=INK, after=2)
    _style(doc, "Heading 1", font="Georgia", size=18, bold=True, color=GREEN, before=14, after=6)
    _style(doc, "Heading 2", font="Georgia", size=13.5, bold=True, color=TERRA, before=12, after=3)
    _style(doc, "TFS Meta", size=9.5, color=MUT, after=4)
    _style(doc, "TFS Speaker", size=11, bold=True, color=INK, before=8, after=0)
    for ls in ("List Bullet",):
        _style(doc, ls, size=11, after=3)

    if os.path.exists(LOGO):
        try:
            doc.add_picture(LOGO, width=Inches(0.9))
        except Exception:
            pass
    doc.add_paragraph(title, style="Title")
    meta = doc.add_paragraph(style="TFS Meta")
    meta.add_run("Example Org meeting recorded " + when + (". " + info if info else ".")).font.size = Pt(9.5)
    _rule(meta)
    if note:
        n = doc.add_paragraph(style="TFS Meta")
        r = n.add_run(note)
        r.italic = True
        r.font.color.rgb = TERRA

    doc.add_heading("Summary", level=1)
    in_actions = False
    for line in summary.splitlines():
        t = line.strip()
        if not t:
            continue
        if t.startswith("## "):
            doc.add_heading(t[3:], level=2)
            in_actions = t[3:].strip().lower().startswith("action")
        elif t.startswith("- ") or t.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            body = t[2:]
            m = re.match(r"^([^:*]{1,40}):\s+(.*)$", body)
            if in_actions and m:
                _runs(p, "**" + m.group(1) + ":** " + m.group(2))
            else:
                _runs(p, body)
        else:
            _runs(doc.add_paragraph(), t)

    if people:
        doc.add_heading("People", level=1)
        pn = doc.add_paragraph(style="TFS Meta")
        pn.add_run("Invited to this meeting. Please check that each email is right. Add a name before an email like this: Jordan Lee: jordan@example.com. "
                   "Delete anyone who was not there.")
        for line in people:
            _runs(doc.add_paragraph(style="List Bullet"), line)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_heading("Transcript", level=1)
    turns = tfsfmt.parse(transcript)
    col = tfsfmt.colors(turns)
    for t in turns:
        if t["who"] or t["time"]:
            p = doc.add_paragraph(style="TFS Speaker")
            r = p.add_run(t["who"] or "Speaker")
            r.bold = True
            r.font.color.rgb = _rgb(col.get(t["who"], "#6b5a4b"))
            if t["time"]:
                r2 = p.add_run("   " + t["time"])
                r2.bold = False
                r2.font.size = Pt(9)
                r2.font.color.rgb = MUT
        doc.add_paragraph(t["text"])

    fp = sec.footer.paragraphs[0]
    fr = fp.add_run("Example Org  |  example.org  |  Private: please do not share")
    fr.font.size = Pt(8)
    fr.font.color.rgb = MUT
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_TIME = re.compile(r"^(.*?)\s{2,}(\d{1,2}:\d{2}(?::\d{2})?)\s*$")


def parse_docx(data: bytes):
    """Return (summary_markdown, transcript_text, people_lines). Any may be empty if the file does not have that part."""
    doc = Document(io.BytesIO(data))
    pars = doc.paragraphs
    tr = next((i for i, p in enumerate(pars) if p.text.strip().lower() == "transcript" and p.style.name.lower().startswith("heading")), None)
    sm = next((i for i, p in enumerate(pars) if p.text.strip().lower() == "summary" and p.style.name.lower().startswith("heading")), None)
    pe = next((i for i, p in enumerate(pars) if p.text.strip().lower() == "people" and p.style.name.lower().startswith("heading")), None)
    summary = ""
    people = []
    if pe is not None:
        for p in pars[pe + 1:(tr if tr is not None else len(pars))]:
            if p.style.name.lower().startswith("list") and p.text.strip():
                people.append(p.text.strip())
    end = len(pars) if tr is None else tr
    if pe is not None and pe < end and (sm is None or pe > sm):
        end = pe
    if sm is not None:
        out = []
        for p in pars[sm + 1:end]:
            t = p.text.strip()
            s = p.style.name
            if not t or s == "TFS Meta":
                continue
            if s.lower().startswith("heading"):
                out.append("## " + t)
            elif s.lower().startswith("list"):
                out.append("- " + t)
            else:
                out.append(t)
        summary = "\n".join(out).strip()
    lines, cur = [], None
    if tr is not None:
        for p in pars[tr + 1:]:
            t = p.text.strip()
            if not t:
                continue
            if p.style.name == "TFS Speaker":
                if cur:
                    lines.append(cur)
                m = _TIME.match(p.text.rstrip())
                who, ts = (m.group(1).strip(), m.group(2)) if m else (t, "")
                cur = ("[" + ts + "] " if ts else "") + who + ": "
            elif cur is None:
                cur = t
            else:
                cur += ("" if cur.endswith(": ") else " ") + t
        if cur:
            lines.append(cur)
    return summary, "\n".join(lines).strip(), people


def norm(s: str) -> str:
    return re.sub(r"[\s*_#\-]+", " ", s).strip().lower()
